import openpyxl, datetime, re
from collections import defaultdict
P='/mnt/user-data/attachments/3ab7db07-7d8f-5b8f-b926-2213c300b28b/0eb5b5ae-_______________________________________.xlsm'
wb=openpyxl.load_workbook(P, data_only=True)
SLOTS=["11:00","12:30","14:00","15:30","17:00","18:30"]; ROOMS=["nam","eli","noy"]
RN={"nam":"Нам","eli":"Ели","noy":"НОЙ"}

def clean(s):
    s=str(s or "")
    note=None
    m=re.search(r'[«"]([^»"]+)[»"]',s)
    if m: note=m.group(1).strip()
    s=re.sub(r'[«"][^»"]*[»"]','',s)
    s=re.sub(r'\s+',' ',s).strip().strip('+').strip()
    return s,note

# клиенты
ws=wb['Учет клиентов Август']
GK=[('Альфа Групп','g_avesto','corporate'),('Бетаи нав','g_koyonot','corporate'),
    ('ГЭМ','g_tgem','corporate'),('Дети','g_deti','child'),('Сертификат','g_sert','certificate'),
    ('Коенот Тичорат','g_tijorat','corporate'),('Клиенты','g_private','private')]
cur=None; clients=[]
for r in ws.iter_rows(values_only=True):
    a,b=r[0],r[1]
    if a and isinstance(a,str):
        for g in GK:
            if a.strip().startswith(g[0]): cur=g; break
    if isinstance(a,(int,float)) and b and isinstance(b,str) and cur:
        name,note=clean(b)
        if not name: continue
        clients.append({"raw":b,"name":name,"note":note,"group":cur[1],"groupTitle":cur[0],
            "seg":cur[2],
            "nam":r[2] or 0,"eli":r[3] or 0,"noy":r[4] or 0,"total":r[5] or 0,
            "start":r[6],"end":r[7],"tel":r[8]})

# индекс: имя -> список карточек
by=defaultdict(list)
for c in clients: by[c["name"].lower()].append(c)

# сетка
ws=wb['График на Август']
grid=[]
for r in ws.iter_rows(min_row=3, values_only=True):
    dt=r[0]
    if not isinstance(dt,datetime.datetime): continue
    cells=list(r[2:20])[:18]
    closed=any(isinstance(c,str) and 'выходной' in c for c in cells if c)
    if closed: continue
    for i,c in enumerate(cells):
        if isinstance(c,str) and c.strip() not in ('','0'):
            name,note=clean(c)
            grid.append({"date":dt.strftime('%Y-%m-%d'),"room":ROOMS[i//6],"slot":SLOTS[i%6],
                         "raw":c,"name":name,"note":note})

print("ячеек с именем в сетке:",len(grid))
# сопоставление
unmatched=[]; ambiguous=[]
for g in grid:
    cand=by.get(g["name"].lower(),[])
    if not cand: unmatched.append(g); continue
    if len(cand)>1:
        pick=[c for c in cand if g["note"] and c["note"]==g["note"]] or [c for c in cand if c["total"]>0]
        if len(pick)!=1: ambiguous.append((g,cand))
        else: g["client"]=pick[0]
    else: g["client"]=cand[0]
print("не нашли карточку:",len(unmatched),"| неоднозначных:",len(ambiguous))
for u in {u["name"] for u in unmatched}: print("  НЕТ КАРТОЧКИ:",repr(u))
for g,c in ambiguous[:5]: print("  НЕОДНОЗНАЧНО:",repr(g["name"]),"->",len(c),"карточек")

# сверка счётчиков карточки с фактической сеткой
fact=defaultdict(lambda: defaultdict(int))
for g in grid:
    if "client" in g: fact[g["client"]["raw"]][g["room"]]+=1
print("\n=== сверка «карточка против сетки» ===")
bad=0
for c in clients:
    if c["total"]==0 and c["raw"] not in fact: continue
    f=fact.get(c["raw"],{})
    card=(c["nam"],c["eli"],c["noy"])
    got=(f.get("nam",0),f.get("eli",0),f.get("noy",0))
    if card!=got:
        bad+=1
        print(f"  {c['name']:26} карточка Нам/Ели/НОЙ={card} сетка={got} итого_карточки={c['total']}")
print("расхождений:",bad,"из",len([c for c in clients if c['total'] or c['raw'] in fact]))
